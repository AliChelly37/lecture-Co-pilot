import io
from datetime import UTC, datetime
from pathlib import Path

from PIL import Image

from lecture_copilot.align import SlideAligner, coverage, off_slide_stretches
from lecture_copilot.deck import SlideEntry, SlideIndex, extract_deck, hotwords_from_index
from lecture_copilot.photos import prepare_photo, seconds_into_lecture


def make_pdf(path: Path, pages: list[str]) -> None:
    import pymupdf

    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        page.insert_text((72, 72), text, fontsize=14)
    doc.save(str(path))
    doc.close()


def make_pptx(path: Path, slides: list[tuple[str, str]]) -> None:
    from pptx import Presentation

    prs = Presentation()
    layout = prs.slide_layouts[1]  # title + content
    for title, body in slides:
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = title
        s.placeholders[1].text = body
    prs.save(str(path))


def test_pdf_and_pptx_extraction(tmp_path: Path) -> None:
    pdf = tmp_path / "deck.pdf"
    make_pdf(pdf, ["Fugacity and the Peng-Robinson equation", "Residual Gibbs energy"])
    d = extract_deck(pdf)
    assert d.filename == "deck.pdf" and len(d.slides) == 2
    assert "Peng-Robinson" in d.slides[0] and "Gibbs" in d.slides[1]
    assert len(d.sha256) == 64

    pptx = tmp_path / "deck.pptx"
    make_pptx(pptx, [("Fugacity", "effective pressure of a real gas"), ("Midterm", "October 21, same room")])
    p = extract_deck(pptx)
    assert len(p.slides) == 2 and "effective pressure" in p.slides[0] and "October 21" in p.slides[1]


def test_aligner_maps_windows_and_finds_off_slide() -> None:
    slides = [
        "Fugacity: effective pressure replacing pressure in chemical potential of a real gas",
        "Peng-Robinson equation of state: compressibility factor Z, cubic in volume",
        "Residual Gibbs energy: integral of (Z - 1)/P dP at constant temperature",
    ]
    aligner = SlideAligner(slides)
    seg = lambda t0, t1, text: {"t0": t0, "t1": t1, "text": text}  # noqa: E731
    segments = [
        seg(0, 25, "today fugacity is the effective pressure for a real gas chemical potential"),
        seg(30, 55, "we use the peng robinson equation of state to get the compressibility factor Z"),
        seg(60, 85, "so last summer I went hiking and my dog got lost in the mountains for three days"),
        seg(90, 115, "anyway the dog came back and we drove home through the rain and stopped for lunch"),
        seg(120, 145, "another story about my neighbour's garden and the tomatoes he grows every summer"),
        seg(150, 175, "okay the tomatoes went into a sauce and we ate it with pasta on the porch"),
        seg(180, 205, "back to the residual gibbs energy which is the integral of Z minus one over P dP"),
    ]
    windows = aligner.align(segments, window_s=30)
    assert [w.slide for w in windows[:2]] == [1, 2]
    assert windows[-1].slide == 3
    assert windows[2].slide is None and windows[3].slide is None
    stretches = off_slide_stretches(windows, min_windows=3)
    assert stretches and stretches[0][0] == 60.0
    assert 0 < coverage(windows) < 1


def test_hotwords_dedup_and_limit() -> None:
    idx = SlideIndex(
        course_terms=["fugacity", "Peng-Robinson"],
        slides=[SlideEntry(slide=1, title="t", key_terms=["fugacity", "residual Gibbs energy"], stated_dates=[], summary="s")],
    )
    assert hotwords_from_index(idx) == "fugacity, Peng-Robinson, residual Gibbs energy"


def test_photo_prepare_downscales_and_reads_exif() -> None:
    img = Image.new("RGB", (4000, 3000), (200, 200, 200))
    exif = Image.Exif()
    exif[36867] = "2026:09:16 10:17:30"  # DateTimeOriginal
    buf = io.BytesIO()
    img.save(buf, format="JPEG", exif=exif.tobytes())
    p = prepare_photo(buf.getvalue(), "board.jpg", "Europe/Paris", max_edge=1280)
    assert max(p.width, p.height) == 1280
    assert p.shot_at is not None and p.shot_at.hour == 10 and p.shot_at.tzinfo is not None
    assert len(p.jpeg) < len(buf.getvalue())

    # 10:17:30 Paris = 08:17:30 UTC; lecture started 08:00 UTC -> 1050 s in.
    started = datetime(2026, 9, 16, 8, 0, tzinfo=UTC).isoformat()
    assert seconds_into_lecture(p.shot_at, started, None) == 1050.0
    ended = datetime(2026, 9, 16, 8, 10, tzinfo=UTC).isoformat()
    assert seconds_into_lecture(p.shot_at, started, ended) is None  # outside window + slack


def test_photo_without_exif_is_unassigned() -> None:
    img = Image.new("RGB", (800, 600), (10, 10, 10))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    p = prepare_photo(buf.getvalue(), "x.png", "UTC")
    assert p.shot_at is None
    assert seconds_into_lecture(None, "2026-09-16T08:00:00+00:00", None) is None
